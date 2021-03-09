import re
import os
import json
import math
import time
import requests
import pymysql
import docx
from win32com import client as wc
from docx import Document
'''
解析拟建名单.docx
'''
data = Document(r'C:\Users\20945\Desktop\Spider\now_spider\高新技术\拟建名单(1).docx')
exit()
all_tables = data.tables
sum = []
exit()
for ta in all_tables:
    rows = ta.rows
    for j in range(1, len(rows)):
        dic = {}
        number = ta.cell(j, 0).text
        center_name = ta.cell(j, 1).text.replace('\n','').partition('市级企业高新技术研究开发中心')[0]
        unit = ta.cell(j, 2).text
        area = ta.cell(j, 3).text

        dic['number'] = number
        dic['center_name'] = center_name
        dic['unit'] = unit
        dic['area'] = area
        print(dic)
        sum.append(dic)

db = pymysql.connect(host="localhost", user="root", password='123456', database="test", port=3306)
cursor = db.cursor()
for i in sum:
    center_name = i.get('center_name')
    supporting_unit = i.get('unit')
    area = i.get('area')

    sql = """insert into high_tech_bak (id,name,support_unit,area,research_center)values(NULL,'{}','{}','{}','市级企业高新技术研究开发中心')""".format(center_name,supporting_unit,area)
    cursor.execute(sql)
    db.commit()
db.close()

